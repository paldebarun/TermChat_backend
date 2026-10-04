from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path

from app.config import get_settings


class UnsupportedDocument(Exception):
    pass


MAX_PDF_PAGES = 500
# A .docx is a zip; refuse ones that inflate far beyond the upload size.
MAX_ZIP_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


def _check_zip_bomb(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile:
        raise ValueError("invalid docx file")
    if total > MAX_ZIP_UNCOMPRESSED_BYTES:
        raise ValueError("document expands beyond the allowed size")


def extract_text(filename: str, content_type: str, data: bytes) -> str:
    settings = get_settings()
    if len(data) > settings.assistant_max_document_bytes:
        raise ValueError("document exceeds configured indexing size limit")

    suffix = Path(filename).suffix.lower()
    ctype = (content_type or "").lower()

    if suffix == ".pdf" or "pdf" in ctype:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("encrypted PDFs are not supported")
        parts: list[str] = []
        total = 0
        for page in reader.pages[:MAX_PDF_PAGES]:
            page_text = page.extract_text() or ""
            parts.append(page_text)
            total += len(page_text)
            if total >= settings.assistant_max_document_chars:
                break  # stop early instead of extracting the whole file
        text = "\n".join(parts)
    elif suffix == ".docx" or "wordprocessingml" in ctype:
        from docx import Document
        _check_zip_bomb(data)
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    elif suffix in {".json", ".jsonl"} or "json" in ctype:
        try:
            obj = json.loads(data.decode("utf-8", errors="replace"))
            text = json.dumps(obj, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            text = data.decode("utf-8", errors="replace")
    elif suffix == ".csv" or "csv" in ctype:
        rows = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
        text = "\n".join(" | ".join(row) for row in rows)
    elif suffix in {".txt", ".md", ".markdown", ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css"} or ctype.startswith("text/"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocument(f"unsupported document type: {content_type or suffix}")

    return text[: settings.assistant_max_document_chars].strip()


def chunk_text(text: str) -> list[tuple[int, str]]:
    settings = get_settings()
    size = max(settings.assistant_chunk_size_chars, 1)  # 0 would never advance
    overlap = min(settings.assistant_chunk_overlap_chars, size // 2)
    if not text:
        return []
    chunks: list[tuple[int, str]] = []
    start = 0
    index = 0
    while start < len(text):
        end = min(start + size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append((index, chunk))
            index += 1
        if end >= len(text):
            break
        start = end - overlap
    return chunks
